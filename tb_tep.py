# -*- coding: utf-8 -*-
"""Xuất file: Word đúng thể thức văn bản hành chính (NĐ 30/2020/NĐ-CP) và Excel bảng đăng ký."""
import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter

from tb_parser import fold
from tb_van_ban import (SHEET_TITLE, _d, bo_cuc, kinh_gui, loi_ket, muc_noi_dung, ngay_chu, tieu_de_trich_yeu)

FONT = 'Times New Roman'
AL = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER, 'right': WD_ALIGN_PARAGRAPH.RIGHT,
      'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}
_PPR_AFTER = ('w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap', 'w:overflowPunct',
              'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN', 'w:bidi', 'w:adjustRightInd', 'w:snapToGrid',
              'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc',
              'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle',
              'w:rPr', 'w:sectPr', 'w:pPrChange')


# ------------------------------------------------------------------ Word
def _run(p, text, size=13, bold=False, italic=False):
    r = p.add_run(text)
    r.font.name, r.font.size, r.bold, r.italic = FONT, Pt(size), bold, italic
    rpr = r._element.get_or_add_rPr()
    rpr.rFonts.set(qn('w:eastAsia'), FONT)
    rpr.rFonts.set(qn('w:cs'), FONT)
    return r


def _p(c, text='', size=13, bold=False, italic=False, align='justify', before=0, after=6, first=None,
       left=None, right=None, line=1.3, reuse=False):
    p = c.paragraphs[0] if reuse else c.add_paragraph()
    f = p.paragraph_format
    f.alignment, f.space_before, f.space_after, f.line_spacing = AL[align], Pt(before), Pt(after), line
    if first is not None:
        f.first_line_indent = Cm(first)
    if left is not None:
        f.left_indent = Cm(left)
    if right is not None:
        f.right_indent = Cm(right)
    if text:
        _run(p, text, size, bold, italic)
    return p


def _gach_chan(p, sz=6):
    """Đường kẻ dưới đoạn (dùng làm gạch chân tên cơ quan, tiêu ngữ, trích yếu)."""
    ppr = p._p.get_or_add_pPr()
    b, bt = OxmlElement('w:pBdr'), OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', str(sz)), ('space', '1'), ('color', '000000')):
        bt.set(qn('w:' + k), v)
    b.append(bt)
    ppr.insert_element_before(b, *_PPR_AFTER)


def _bang(doc, cols, widths):
    t = doc.add_table(rows=1, cols=cols)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, w in zip(t.rows[0].cells, widths):
        c.width = Cm(w)
    return t


def _them_hang(t, widths):
    row = t.add_row()
    for c, w in zip(row.cells, widths):
        c.width = Cm(w)
    return row.cells


def tao_docx(d):
    """Thông báo đúng thể thức: A4, lề trái 3cm – phải 2cm – trên/dưới 2cm, Times New Roman 13–14."""
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(3), Cm(2), Cm(2), Cm(2)
    st = doc.styles['Normal']
    st.font.name, st.font.size = FONT, Pt(13)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

    # Quốc hiệu – tiêu ngữ | cơ quan ban hành
    W = [7.0, 9.0]
    t = _bang(doc, 2, W)
    a, b = t.rows[0].cells
    _p(a, d['don_vi_chu_quan'], 12, align='center', after=0, line=1.0, reuse=True)
    p = _p(a, d['don_vi'], 12, bold=True, align='center', after=0, line=1.0, left=1.2, right=1.2)
    _gach_chan(p)
    _p(b, 'CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM', 12, bold=True, align='center', after=0, line=1.0, reuse=True)
    p = _p(b, 'Độc lập – Tự do – Hạnh phúc', 13, bold=True, align='center', after=0, line=1.0, left=1.8, right=1.8)
    _gach_chan(p)
    a, b = _them_hang(t, W)
    _p(a, 'Số: ' + d['so_vb'], 13, align='center', before=8, after=0, line=1.0, reuse=True)
    _p(b, '%s, %s' % (d['dia_danh'], ngay_chu(_d(d['ngay_ban_hanh']))), 13, italic=True, align='center', before=8,
       after=0, line=1.0, reuse=True)

    # Tên loại + trích yếu
    _p(doc, 'THÔNG BÁO', 14, bold=True, align='center', before=16, after=2, line=1.0)
    p = _p(doc, tieu_de_trich_yeu(d), 13, bold=True, align='center', after=0, line=1.15, left=1.5, right=1.5)
    g = _p(doc, '', 4, align='center', after=10, line=1.0, left=6.0, right=6.0)
    _gach_chan(g)

    kg = kinh_gui(d)
    if kg:
        p = _p(doc, '', align='left', after=6, first=1.0)
        _run(p, 'Kính gửi: ', 13, bold=True)
        _run(p, kg, 13)
    _p(doc, '%s thông báo đến sinh viên về việc %s %s như sau:' % (d['don_vi_viet'], d['muc_dich'], d['ten_su_kien']),
       first=1.0)

    for i, (nhan, nd) in enumerate(muc_noi_dung(d), 1):
        p = _p(doc, '', first=1.0)
        _run(p, '%d. %s: ' % (i, nhan), 13, bold=True)
        _run(p, nd, 13)
        if nhan == 'Chỉ tiêu':
            CW = [3.0, 3.5, 8.5]
            tb = _bang(doc, 3, CW)
            tb.style = 'Table Grid'
            for c, txt, al in zip(tb.rows[0].cells, ['Khóa', 'Số lượng', 'Ghi chú'], ['center'] * 3):
                _p(c, txt, 13, bold=True, align=al, after=2, before=2, line=1.0, reuse=True)
            for it in d['chi_tieu']:
                cs = _them_hang(tb, CW)
                _p(cs[0], it['khoa'], 13, align='center', after=2, before=2, line=1.0, reuse=True)
                _p(cs[1], str(it['so_luong']), 13, align='center', after=2, before=2, line=1.0, reuse=True)
                _p(cs[2], it['ghi_chu'], 13, align='left', after=2, before=2, line=1.0, reuse=True)
            cs = _them_hang(tb, CW)
            _p(cs[0], 'Cộng', 13, bold=True, align='center', after=2, before=2, line=1.0, reuse=True)
            _p(cs[1], str(d['tong_chi_tieu']), 13, bold=True, align='center', after=2, before=2, line=1.0, reuse=True)
            _p(cs[2], '', 13, after=2, before=2, line=1.0, reuse=True)
            _p(doc, '', 6, after=4, line=1.0)

    _p(doc, loi_ket(d),
       first=1.0, before=4, after=10)

    # Nơi nhận | chức vụ người ký
    FW = [8.0, 8.0]
    ft = _bang(doc, 2, FW)
    a, b = ft.rows[0].cells
    _p(a, 'Nơi nhận:', 12, bold=True, italic=True, align='left', after=0, line=1.0, reuse=True)
    _p(a, '- Như trên;', 11, align='left', after=0, line=1.0)
    _p(a, '- Lưu: VT.', 11, align='left', after=0, line=1.0)
    _p(b, d['chuc_danh_ky'], 13, bold=True, align='center', after=0, line=1.0, reuse=True)
    _p(b, '', 13, after=0, line=1.0)
    _p(b, '', 13, after=0, line=1.0)
    _p(b, '', 13, after=0, line=1.0)
    _p(b, d['nguoi_ky'], 13, bold=True, align='center', after=0, line=1.0)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ Excel (bảng đăng ký)
NAVY, BLUE, PALE = '1E3A8A', '2563EB', 'DCE6FA'
_THIN = Side(style='thin', color='CBD5E1')
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def tao_xlsx(d):
    L = bo_cuc(d)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE
    base = Font(name='Arial', size=11)
    for col, w in zip('ABCD', (7, 9, 20, 40)):
        ws.column_dimensions[col].width = w
    for r, txt, fill, fnt, h in (
            (1, L['title'], NAVY, Font(name='Arial', size=14, bold=True, color='FFFFFF'), 40),
            (2, L['info'], 'EEF2F9', Font(name='Arial', size=11, color='0F172A'), 18 * L['info_lines'] + 8),
            (3, L['guide'], 'FFF7E0', Font(name='Arial', size=10, italic=True, color='7A4B00'), 46)):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        c = ws.cell(r, 1, txt)
        c.font, c.fill = fnt, PatternFill('solid', fgColor=fill)
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center' if r == 1 else 'left', indent=0 if r == 1 else 1)
        ws.row_dimensions[r].height = h
    for i, h in enumerate(L['header'], 1):
        c = ws.cell(4, i, h)
        c.font, c.fill = Font(name='Arial', size=11, bold=True, color='FFFFFF'), PatternFill('solid', fgColor=BLUE)
        c.alignment, c.border = Alignment(horizontal='center', vertical='center'), BORDER
    ws.row_dimensions[4].height = 26
    for it in L['rows']:
        r = it['row']
        if it['kind'] == 'group':
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
            c = ws.cell(r, 1, it['label'])
            c.font, c.fill = Font(name='Arial', size=11, bold=True, color=NAVY), PatternFill('solid', fgColor=PALE)
            c.alignment = Alignment(vertical='center', indent=1, wrap_text=True)
            for k in range(1, 5):
                ws.cell(r, k).border = BORDER
            ws.row_dimensions[r].height = 24
        else:
            vals = [it['stt'], it['khoa'], None, None]
            for k, v in enumerate(vals, 1):
                c = ws.cell(r, k, v)
                c.font, c.border = base, BORDER
                c.alignment = Alignment(horizontal='center' if k <= 2 else 'left', vertical='center')
                if k >= 3:
                    c.protection = Protection(locked=False)
                    c.number_format = '@'
            ws.row_dimensions[r].height = 22
    a, z = L['first'], L['last']
    ws.conditional_formatting.add('C%d:C%d' % (a, z), FormulaRule(
        formula=['AND(LEN($C%d)>0,COUNTIF($C$%d:$C$%d,$C%d)>1)' % (a, a, z, a)],
        fill=PatternFill('solid', bgColor='FECACA', fgColor='FECACA'), font=Font(color='B91C1C', bold=True)))
    ws.freeze_panes = 'A5'
    ws.print_title_rows = '4:4'
    ws.page_setup.orientation, ws.page_setup.paperSize, ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 'portrait', 9, 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.protection.sheet = True          # khóa các ô tiêu đề; chỉ MSSV + Họ tên sửa được (không đặt mật khẩu)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def tao_xlsx_da_dang_ky(entry, rows, tra_cuu=None):
    """Xuất danh sách sinh viên đã điền vào Google Sheet.
    entry: bản ghi lịch sử (ten, chi_tieu…); rows: [{'khoa','mssv','ho_ten'}]
    tra_cuu: hàm tùy chọn mssv -> {'ho_ten':…, 'khoa_hoc':…} hoặc None (để đối chiếu với CSDL sinh viên)."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Đã đăng ký'
    heads = ['STT', 'Khóa', 'MSSV', 'Họ và tên', 'Đối chiếu']
    for i, h in enumerate(heads, 1):
        c = ws.cell(1, i, h)
        c.font, c.fill = Font(name='Arial', bold=True, color='FFFFFF'), PatternFill('solid', fgColor=NAVY)
        c.alignment, c.border = Alignment(horizontal='center'), BORDER
    dem_mssv = {}
    for r in rows:
        if r['mssv']:
            dem_mssv[r['mssv']] = dem_mssv.get(r['mssv'], 0) + 1
    thuc = {}
    n = 0
    for r in rows:
        if not (r['mssv'] or r['ho_ten']):
            continue
        n += 1
        thuc[r['khoa']] = thuc.get(r['khoa'], 0) + 1
        ghi = []
        if not r['mssv'] or not r['ho_ten']:
            ghi.append('Thiếu MSSV/họ tên')
        if r['mssv'] and dem_mssv[r['mssv']] > 1:
            ghi.append('MSSV trùng')
        if tra_cuu and r['mssv']:
            sv = tra_cuu(r['mssv'])
            if not sv:
                ghi.append('Không có trong danh sách sinh viên')
            else:
                if fold(sv.get('ho_ten', '')).split() != fold(r['ho_ten']).split():
                    ghi.append('Họ tên khác hồ sơ: ' + sv.get('ho_ten', ''))
                kh = ''.join(ch for ch in str(sv.get('khoa_hoc', '')) if ch.isdigit())
                if kh and kh != str(r['khoa']):
                    ghi.append('Sinh viên thuộc khóa %s' % kh)
        for i, v in enumerate([n, r['khoa'], r['mssv'], r['ho_ten'], '; '.join(ghi)], 1):
            c = ws.cell(n + 1, i, v)
            c.font, c.border = Font(name='Arial', size=11, color='B91C1C' if (i == 5 and ghi) else '000000'), BORDER
    for col, w in zip('ABCDE', (7, 9, 18, 36, 46)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    s2 = wb.create_sheet('Tổng hợp')
    for i, h in enumerate(['Khóa', 'Chỉ tiêu', 'Đã đăng ký', 'Còn lại'], 1):
        c = s2.cell(1, i, h)
        c.font, c.fill = Font(name='Arial', bold=True, color='FFFFFF'), PatternFill('solid', fgColor=NAVY)
    for j, it in enumerate(entry.get('chi_tieu', []), 2):
        s2.cell(j, 1, 'Khóa ' + str(it['khoa']))
        s2.cell(j, 2, it['so_luong'])
        s2.cell(j, 3, thuc.get(str(it['khoa']), 0))
        s2.cell(j, 4, '=B%d-C%d' % (j, j))
    for col, w in zip('ABCD', (14, 12, 14, 12)):
        s2.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()