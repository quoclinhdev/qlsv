# -*- coding: utf-8 -*-
"""Quản lý sinh viên – Flask + SQLite.
Chạy:  python app.py   (mặc định http://127.0.0.1:5000)
"""
import io
import os
import secrets
import sys
import unicodedata
from functools import wraps
from pathlib import Path

from flask import (Flask, abort, flash, g, jsonify, redirect, render_template, request,
                   send_file, session, url_for)

import csdl
import routes_chamdiem
from csdl import HOC_KY, KET_QUA

try:  # console Windows (cp1252) không in được tiếng Việt
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

BASE = Path(__file__).parent
app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024  # file Excel tối đa 10 MB
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"


def _secret():
    if os.environ.get("SECRET_KEY"):
        return os.environ["SECRET_KEY"]
    f = BASE / "secret.key"
    if not f.exists():
        f.write_text(secrets.token_hex(32))
    return f.read_text().strip()


app.secret_key = _secret()


# ------------------------------------------------------------------ hạ tầng
def db():
    if "db" not in g:
        g.db = csdl.DB()
    return g.db


@app.teardown_appcontext
def _close(_e):
    d = g.pop("db", None)
    if d:
        d.close()


def csrf_token():
    if "csrf" not in session:
        session["csrf"] = secrets.token_hex(16)
    return session["csrf"]


app.jinja_env.globals["csrf_token"] = csrf_token
app.jinja_env.globals.update(HOC_KY=HOC_KY, KET_QUA=KET_QUA)


@app.before_request
def _csrf_check():
    if request.method == "POST":
        tok = request.form.get("csrf") or request.headers.get("X-CSRF")
        if not tok or tok != session.get("csrf"):
            abort(400, "CSRF token không hợp lệ – hãy tải lại trang.")


def login_required(f):  # không cần đăng nhập: giữ tên để các route không phải sửa
    return f


admin_required = login_required


def gui_excel(ten, hoc, rows_title, headers, rows, widths, phu_luc=None):
    bio = io.BytesIO()
    csdl.xuat_excel(bio, rows_title, headers, rows, widths, phu_luc)
    bio.seek(0)
    return send_file(bio, as_attachment=True, download_name=ten,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _nam(v, mac_dinh):
    """Năm học hợp lệ dạng 2026-2027, nếu sai thì dùng giá trị mặc định."""
    v = (v or "").strip()
    return v if csdl.nam_hoc_hop_le(v) else mac_dinh


def _ds_mssv(j):
    ds = j.get("mssvs") if isinstance(j, dict) else None
    if not isinstance(ds, list) or not ds:
        abort(400, "Chưa chọn sinh viên nào.")
    return [str(x).strip() for x in ds if str(x).strip()]


def so(v, kind=float):
    v = (v or "").strip().replace(",", ".")
    return kind(float(v)) if v else None


# ------------------------------------------------------------------ sinh viên
@app.route("/")
@login_required
def home():
    return redirect(url_for("sinh_vien"))


@app.route("/sinh-vien")
@login_required
def sinh_vien():
    # Lọc/tìm kiếm/sắp xếp/phân trang do DataTables xử lý phía trình duyệt
    return render_template("sinh_vien.html", rows=db().ds_sinh_vien())


@app.post("/sinh-vien/xoa-nhieu")
@admin_required
def sv_xoa_nhieu():
    n = db().xoa_sinh_vien_nhieu(_ds_mssv(request.get_json(silent=True)))
    return jsonify(da_xoa=n)


SV_FIELDS = ["mssv", "ho_lot", "ten", "lop", "ma_lop", "khoa_hoc", "nganh", "khoa"]


@app.route("/sinh-vien/them", methods=["GET", "POST"])
@admin_required
def sv_them():
    if request.method == "POST":
        d = {k: request.form.get(k, "").strip() for k in SV_FIELDS}
        if not d["mssv"]:
            flash("MSSV không được trống.", "err")
        elif db().con.execute("SELECT 1 FROM sinh_vien WHERE mssv=?", (d["mssv"],)).fetchone():
            flash("MSSV đã tồn tại.", "err")
        else:
            db().luu_sinh_vien(d)
            flash("Đã thêm sinh viên.", "ok")
            return redirect(url_for("sinh_vien"))
        return render_template("sv_form.html", sv=d, moi=True)
    return render_template("sv_form.html", sv={}, moi=True)


@app.route("/sinh-vien/<mssv>/sua", methods=["GET", "POST"])
@admin_required
def sv_sua(mssv):
    sv = db().con.execute("SELECT * FROM sinh_vien WHERE mssv=?", (mssv,)).fetchone()
    if not sv:
        abort(404)
    if request.method == "POST":
        d = {k: request.form.get(k, "").strip() for k in SV_FIELDS}
        d["mssv"] = mssv
        db().sua_sinh_vien(d)
        flash("Đã cập nhật.", "ok")
        return redirect(url_for("sinh_vien"))
    return render_template("sv_form.html", sv=sv, moi=False)


@app.post("/sinh-vien/<mssv>/xoa")
@admin_required
def sv_xoa(mssv):
    db().xoa_sinh_vien(mssv)
    flash(f"Đã xóa sinh viên {mssv} cùng điểm rèn luyện, tình trạng và điểm danh.", "ok")
    return redirect(url_for("sinh_vien"))


SV_XUAT_COT = [("STT", 6), ("MSSV", 14), ("Họ lót", 24), ("Tên", 12), ("Lớp", 14), ("Mã lớp", 14),
               ("Khóa", 8), ("Ngành", 30), ("Khoa", 40)]


def _xuat_sv(rows, ten_file):
    data = [[i, *[(r[k] if r[k] is not None else "") for k in
                  ("mssv", "ho_lot", "ten", "lop", "ma_lop", "khoa_hoc", "nganh", "khoa")]]
            for i, r in enumerate(rows, 1)]
    return gui_excel(ten_file, None, f"DANH SÁCH SINH VIÊN ({len(data)} sinh viên)",
                     [c for c, _ in SV_XUAT_COT], data, [w for _, w in SV_XUAT_COT])


@app.get("/sinh-vien/xuat")
@login_required
def sv_xuat():
    """Xuất Excel theo đúng bộ lọc đang áp dụng trên trang (mã lớp, khóa, khoa, ô tìm kiếm)."""
    lop, kh, khoa, q = (request.args.get(k, "").strip() for k in ("lop", "kh", "khoa", "q"))
    qn = _bo_dau(q)
    rows = []
    for r in db().ds_sinh_vien():
        if lop and (r["ma_lop"] or "").strip() != lop:
            continue
        if kh and str(r["khoa_hoc"] or "").strip() != kh:
            continue
        if khoa and (r["khoa"] or "").strip() != khoa:
            continue
        if qn and qn not in _bo_dau(" ".join(str(r[k] or "") for k in
                                            ("mssv", "ho_lot", "ten", "lop", "ma_lop", "khoa_hoc", "nganh", "khoa"))):
            continue
        rows.append(r)
    return _xuat_sv(rows, "SinhVien.xlsx")


@app.post("/sinh-vien/xuat-chon")
@login_required
def sv_xuat_chon():
    """Xuất Excel các sinh viên đang được tick chọn."""
    ds = {x.strip() for x in request.form.get("mssvs", "").split(",") if x.strip()}
    if not ds:
        abort(400, "Chưa chọn sinh viên nào.")
    return _xuat_sv([r for r in db().ds_sinh_vien() if str(r["mssv"]) in ds], "SinhVien_DaChon.xlsx")


# ---- thêm / cập nhật sinh viên hàng loạt bằng Excel
SV_KHOA_HOC = {"khóa", "khoá", "khóa học", "khoá học", "khoa hoc"}  # "Khóa" (K24) khác "Khoa" (đơn vị)
SV_TIEU_DE = {"mssv": "mssv", "ma sv": "mssv", "ma sinh vien": "mssv", "ma so sinh vien": "mssv",
              "ho lot": "ho_lot", "ho va ten lot": "ho_lot", "ten": "ten", "ho ten": "ho_ten", "ho va ten": "ho_ten",
              "lop": "lop", "ma lop": "ma_lop", "nganh": "nganh", "khoa": "khoa"}


def _cot_sv(h):
    raw = unicodedata.normalize("NFC", " ".join(str(h or "").split()).lower())
    return "khoa_hoc" if raw in SV_KHOA_HOC else SV_TIEU_DE.get(_bo_dau(h))


def _o_chu(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def doc_sv_excel(fobj):
    """Đọc file Excel danh sách sinh viên. Trả về [(số dòng Excel, dict các trường)].
    Tiêu đề nhận: MSSV, Họ lót, Tên (hoặc Họ tên), Lớp, Mã lớp, Khóa, Ngành, Khoa (cột khác như STT được bỏ qua)."""
    from openpyxl import load_workbook
    wb = load_workbook(fobj, read_only=True, data_only=True)
    try:
        ws = wb.active
        cot, kq = None, []
        for idx, row in enumerate(ws.iter_rows(values_only=True), 1):
            if cot is None:
                if idx > 15:
                    break
                m = {}
                for i, h in enumerate(row):
                    k = _cot_sv(h)
                    if k and k not in m:
                        m[k] = i
                if "mssv" in m and ("ten" in m or "ho_ten" in m):
                    cot = m
                continue
            d = {k: (_o_mssv(row[i]) if k == "mssv" else _o_chu(row[i])) if i < len(row) else ""
                 for k, i in cot.items()}
            if not any(d.values()):
                continue
            if "ho_ten" in d:
                if not d.get("ten") and d["ho_ten"]:
                    ho, _, ten = d["ho_ten"].rpartition(" ")
                    d["ho_lot"], d["ten"] = ho, ten
                del d["ho_ten"]
            kq.append((idx, d))
        if cot is None:
            raise ValueError("không tìm thấy dòng tiêu đề có cột MSSV và Tên (hoặc Họ tên)")
        return kq
    finally:
        wb.close()


@app.post("/sinh-vien/nhap")
@admin_required
def sv_nhap():
    f = request.files.get("file")
    cap_nhat = request.form.get("cap_nhat") == "1"
    if not f or not f.filename.lower().endswith(".xlsx"):
        flash("Hãy chọn file .xlsx có cột MSSV, Họ lót, Tên.", "err")
        return redirect(url_for("sinh_vien"))
    try:
        ds = doc_sv_excel(io.BytesIO(f.read()))
    except Exception as e:
        flash(f"Lỗi đọc file: {e}", "err")
        return redirect(url_for("sinh_vien"))
    d_b = db()
    co = {r["mssv"]: r for r in d_b.con.execute("SELECT * FROM sinh_vien")}
    them = sua = 0
    bo_qua, loi, trung, thay = [], [], [], set()
    try:
        for dong, d in ds:
            m = d.get("mssv", "")
            if not m:
                loi.append(f"dòng {dong}: thiếu MSSV")
                continue
            if m not in co and not d.get("ten") and not d.get("ho_lot"):  # sinh viên mới bắt buộc có họ tên
                loi.append(f"dòng {dong}: {m} thiếu họ tên")
                continue
            if m in thay:
                trung.append(m)
                continue
            thay.add(m)
            if m in co:
                if not cap_nhat:
                    bo_qua.append(m)
                    continue
                moi = {k: (d.get(k) or (co[m][k] or "")) for k in SV_FIELDS}  # chỉ ghi đè ô có dữ liệu
                moi["mssv"] = m
                d_b.sua_sinh_vien(moi)
                sua += 1
            else:
                moi = {k: d.get(k, "") for k in SV_FIELDS}
                moi["mssv"] = m
                d_b.luu_sinh_vien(moi)
                them += 1
    except Exception as e:
        flash(f"Dừng giữa chừng do lỗi: {e}. Đã thêm {them}, cập nhật {sua} trước khi lỗi.", "err")
        return redirect(url_for("sinh_vien"))
    msg = f"Đã thêm {them} sinh viên mới"
    if cap_nhat:
        msg += f", cập nhật {sua} sinh viên"
    msg += "."
    if bo_qua:
        msg += (f" Bỏ qua {len(bo_qua)} MSSV đã tồn tại (tích 'Cập nhật' nếu muốn ghi đè): "
                + ", ".join(bo_qua[:8]) + ("…" if len(bo_qua) > 8 else "") + ".")
    if trung:
        msg += f" {len(trung)} MSSV lặp trong file (chỉ lấy dòng đầu): " + ", ".join(trung[:8]) + ("…" if len(trung) > 8 else "") + "."
    if loi:
        msg += f" {len(loi)} dòng lỗi: " + "; ".join(loi[:6]) + ("…" if len(loi) > 6 else "") + "."
    flash(msg, "ok" if (them or sua) else "err")
    return redirect(url_for("sinh_vien"))


@app.get("/sinh-vien/mau")
@login_required
def sv_mau():
    """File Excel mẫu để nhập sinh viên."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Sinh viên"
    tieu_de = ["MSSV", "Họ lót", "Tên", "Lớp", "Mã lớp", "Khóa", "Ngành", "Khoa"]
    ws.append(tieu_de)
    for i, (c, w) in enumerate(zip(tieu_de, [16, 24, 12, 16, 14, 8, 30, 40]), 1):
        cell = ws.cell(1, i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = w
        for r in range(2, 502):  # định dạng văn bản để không mất số 0 đầu MSSV
            ws.cell(r, i).number_format = "@"
    ws.freeze_panes = "A2"
    hd = wb.create_sheet("Hướng dẫn")
    for dong in ["HƯỚNG DẪN NHẬP SINH VIÊN BẰNG EXCEL",
                 "1. Điền dữ liệu từ dòng 2 của sheet 'Sinh viên', mỗi dòng một sinh viên.",
                 "2. Bắt buộc có MSSV và họ tên (Họ lót + Tên). Các cột còn lại có thể để trống.",
                 "3. Có thể dùng một cột 'Họ tên' thay cho 'Họ lót' và 'Tên' (hệ thống tự tách tên là từ cuối).",
                 "4. Cột 'Khóa' là khóa học (vd 24, 25); cột 'Khoa' là đơn vị (vd Khoa Công nghệ thông tin).",
                 "5. MSSV đã có trong hệ thống sẽ bị bỏ qua, trừ khi tích 'Cập nhật nếu MSSV đã tồn tại'.",
                 "6. Có thể nhập lại trực tiếp file đã xuất từ chức năng 'Xuất Excel' (cột STT được bỏ qua)."]:
        hd.append([dong])
    hd["A1"].font = Font(bold=True, size=13)
    hd.column_dimensions["A"].width = 110
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return send_file(bio, as_attachment=True, download_name="MauNhapSinhVien.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ------------------------------------------------------------------ điểm rèn luyện
def _drl_args():
    hk0, nam0 = db().hoc_ky_gan_nhat()  # cùng kỳ mặc định với module Tình trạng
    hk = request.values.get("hk", hk0)
    return (hk if hk in HOC_KY else hk0), _nam(request.values.get("nam"), nam0), request.values.get("lop", "")


@app.route("/diem-ren-luyen")
@login_required
def drl():
    hk, nam, lop = _drl_args()
    rows = db().ds_drl(hk, nam, lop)
    return render_template("drl.html", rows=rows, hk=hk, nam=nam, lop=lop, lops=db().ds_lop(),
                           nams=db().ds_nam_hoc(nam), quy_tac=csdl.QUY_TAC_DRL)


@app.post("/api/drl")
@login_required
def api_drl():
    j = request.get_json(silent=True) or {}
    try:
        diem = csdl.parse_diem(j.get("diem", ""))
    except ValueError as e:
        return jsonify(error=str(e)), 400
    if j.get("hoc_ky") not in HOC_KY or not csdl.nam_hoc_hop_le(j.get("nam_hoc")):
        return jsonify(error="Học kỳ / năm học không hợp lệ"), 400
    if not db().con.execute("SELECT 1 FROM sinh_vien WHERE mssv=?", (j.get("mssv"),)).fetchone():
        return jsonify(error="Không tìm thấy MSSV"), 404
    dac_biet = diem is not None and float(diem) in (1.0, 2.0, 3.0)
    ly = str(j.get("ly_do") or "").strip()[:500] if "ly_do" in j else None
    if dac_biet and ly is not None and not ly:
        return jsonify(error="Cần nhập lý do cho trường hợp Nghỉ / Bảo lưu / Không nộp."), 400
    db().luu_drl(j["mssv"], j["hoc_ky"], j["nam_hoc"].strip(), diem, ly if dac_biet else None)
    return jsonify(diem=diem, xep_loai=csdl.xep_loai_drl(diem),
                   ly_do=db().ly_do_drl(j["mssv"], j["hoc_ky"], j["nam_hoc"].strip()))


@app.post("/api/drl/xoa")
@login_required
def api_drl_xoa():
    j = request.get_json(silent=True) or {}
    if j.get("hoc_ky") not in HOC_KY or not csdl.nam_hoc_hop_le(j.get("nam_hoc")):
        return jsonify(error="Học kỳ / năm học không hợp lệ"), 400
    if not isinstance(j.get("mssvs"), list) or not j["mssvs"]:
        return jsonify(error="Chưa chọn sinh viên nào."), 400
    n = db().xoa_drl_nhieu(j["mssvs"], j["hoc_ky"], j["nam_hoc"].strip())
    return jsonify(da_xoa=n)


def _xuat_drl(ds_rows, hk, nam):
    rows = [[i, r["mssv"], r["ho_lot"], r["ten"], r["diem"], r["xep_loai"] or "", r["lop"], r["ma_lop"], r["ly_do"] or ""]
            for i, r in enumerate(ds_rows, 1)]
    pl = [("Điểm", "Xếp loại")] + [(f"≥ {n}", t) for n, t in csdl.QUY_TAC_DRL[::-1]] + \
         [("1", "Nghỉ"), ("2", "Bảo lưu"), ("3", "Không nộp")]
    return gui_excel(f"DRL_{hk}_{nam}.xlsx", None, f"ĐIỂM RÈN LUYỆN {hk} NĂM HỌC {nam}",
                     ["STT", "MSSV", "HỌ", "TÊN", "ĐIỂM", "XẾP LOẠI", "LỚP", "MÃ LỚP", "LÝ DO"], rows,
                     [6, 14, 24, 12, 8, 12, 24, 14, 60], pl)


@app.get("/diem-ren-luyen/xuat")
@login_required
def drl_xuat():
    hk, nam, lop = _drl_args()
    return _xuat_drl(db().ds_drl(hk, nam, lop), hk, nam)


@app.post("/diem-ren-luyen/xuat-chon")
@login_required
def drl_xuat_chon():
    """Xuất Excel điểm rèn luyện của các sinh viên đang được tick chọn."""
    hk, nam, _lop = _drl_args()
    ds = {x.strip() for x in request.form.get("mssvs", "").split(",") if x.strip()}
    if not ds:
        abort(400, "Chưa chọn sinh viên nào.")
    return _xuat_drl([r for r in db().ds_drl(hk, nam, "") if str(r["mssv"]) in ds], hk, nam)


@app.post("/api/drl/cap-nhat-nhieu")
@login_required
def api_drl_cap_nhat_nhieu():
    """Cập nhật điểm hàng loạt: che_do='cong' (cộng/trừ) hoặc 'dat' (đặt bằng)."""
    j = request.get_json(silent=True) or {}
    if j.get("hoc_ky") not in HOC_KY or not csdl.nam_hoc_hop_le(j.get("nam_hoc")):
        return jsonify(error="Học kỳ / năm học không hợp lệ"), 400
    if not isinstance(j.get("mssvs"), list) or not j["mssvs"]:
        return jsonify(error="Chưa chọn sinh viên nào."), 400
    che_do = j.get("che_do")
    try:
        if che_do == "cong":
            gt = float(str(j.get("gia_tri", "")).strip().replace(",", "."))
            if gt != gt or abs(gt) > 100 or gt == 0:
                raise ValueError("Số điểm cộng/trừ phải khác 0 và nằm trong khoảng -100 đến 100")
        elif che_do == "dat":
            gt = csdl.parse_diem(j.get("gia_tri", ""))
            if gt is None:
                raise ValueError("Hãy nhập điểm cần đặt")
        else:
            raise ValueError("Chế độ cập nhật không hợp lệ")
    except ValueError as e:
        return jsonify(error=str(e) if "điểm" in str(e).lower() or "Điểm" in str(e) else "Giá trị không hợp lệ"), 400
    ly = str(j.get("ly_do") or "").strip()[:500]
    if che_do == "dat" and float(gt) in (1.0, 2.0, 3.0) and not ly:
        return jsonify(error="Cần nhập lý do khi đặt Nghỉ / Bảo lưu / Không nộp."), 400
    return jsonify(db().cap_nhat_drl_nhieu(j["mssvs"], j["hoc_ky"], j["nam_hoc"].strip(), che_do, gt, ly or None))


@app.post("/diem-ren-luyen/nhap")
@login_required
def drl_nhap():
    hk, nam, lop = _drl_args()
    f = request.files.get("file")
    if not f or not f.filename.lower().endswith(".xlsx"):
        flash("Hãy chọn file .xlsx (cột MSSV và ĐIỂM, hoặc MSSV + Ghi chú + Lý do).", "err")
    else:
        try:
            ok, miss, thieu = db().nhap_drl_excel(io.BytesIO(f.read()), hk, nam)
            flash(f"Đã ghi {ok} điểm. Không hợp lệ/không có trong CSDL: {len(miss)}"
                  + (f" ({', '.join(miss[:10])})" if miss else "")
                  + (f". Còn {len(thieu)} sinh viên Nghỉ/Bảo lưu/Không nộp chưa có lý do." if thieu else ""), "ok")
        except Exception as e:  # file sai định dạng
            flash(f"Lỗi nhập file: {e}", "err")
    return redirect(url_for("drl", hk=hk, nam=nam, lop=lop))


# ------------------------------------------------------------------ tình trạng sinh viên
def _tt_args():
    hk0, nam0 = db().hoc_ky_gan_nhat()
    hk = request.values.get("hk", hk0)
    return (hk if hk in HOC_KY else hk0), _nam(request.values.get("nam"), nam0)


@app.route("/tinh-trang")
@login_required
def tinh_trang():
    hk, nam = _tt_args()
    khoa = request.args.get("khoa", "")  # lọc theo khoa ở server; kết quả/khóa/tìm kiếm do DataTables lọc
    rows = db().ds_tinh_trang(hk, nam, khoa=khoa)
    dem = {k: sum(1 for r in rows if r["ket_qua"] == k) for k in KET_QUA}
    return render_template("tinh_trang.html", rows=rows, hk=hk, nam=nam, khoa=khoa, dem=dem,
                           nams=db().ds_nam_hoc(nam), khoas=db().ds_khoa(),
                           svs=db().ds_sinh_vien())


@app.post("/tinh-trang/xoa-nhieu")
@admin_required
def tt_xoa_nhieu():
    j = request.get_json(silent=True) or {}
    if j.get("hoc_ky") not in HOC_KY or not csdl.nam_hoc_hop_le(j.get("nam_hoc")):
        return jsonify(error="Học kỳ / năm học không hợp lệ"), 400
    n = db().xoa_tinh_trang_nhieu(_ds_mssv(j), j["hoc_ky"], j["nam_hoc"].strip())
    return jsonify(da_xoa=n)


@app.route("/tinh-trang/<mssv>/sua", methods=["GET", "POST"])
@admin_required
def tt_sua(mssv):
    hk, nam = _tt_args()
    sv = db().con.execute("SELECT * FROM sinh_vien WHERE mssv=?", (mssv,)).fetchone()
    if not sv:
        abort(404)
    cur = db().con.execute("SELECT * FROM tinh_trang WHERE mssv=? AND hoc_ky=? AND nam_hoc=?",
                           (mssv, hk, nam)).fetchone()
    if request.method == "POST":
        f = request.form
        if f.get("ket_qua") not in KET_QUA:
            abort(400)
        try:
            d = dict(mssv=mssv, hoc_ky=hk, nam_hoc=nam, ket_qua=f["ket_qua"],
                     dtb_hk=so(f.get("dtb_hk")), stc_hk=so(f.get("stc_hk"), int), xl_hk=f.get("xl_hk", "").strip(),
                     no_hk=so(f.get("no_hk"), int), dtb_tl=so(f.get("dtb_tl")), stc_tl=so(f.get("stc_tl"), int),
                     xl_tl=f.get("xl_tl", "").strip(), no_tl=so(f.get("no_tl"), int),
                     ly_do=f.get("ly_do", "").strip())
        except ValueError:
            flash("Các ô số phải nhập đúng định dạng số.", "err")
            return render_template("tt_form.html", sv=sv, t=f, hk=hk, nam=nam)
        db().luu_tinh_trang(d)
        flash("Đã lưu tình trạng và lý do.", "ok")
        return redirect(url_for("tinh_trang", hk=hk, nam=nam))
    return render_template("tt_form.html", sv=sv, t=cur or {"ket_qua": "Bình thường"}, hk=hk, nam=nam)


@app.post("/tinh-trang/nhap")
@admin_required
def tt_nhap():
    hk, nam = _tt_args()
    f = request.files.get("file")
    if not f or not f.filename.lower().endswith(".xlsx"):
        flash("Hãy chọn file .xlsx.", "err")
    else:
        try:
            t = db().nhap_tien_do_excel(io.BytesIO(f.read()), hk, nam)
            flash(f"Đã nhập {t['sinh_vien']} sinh viên xét tiến độ: Bình thường {t['binh_thuong']}, "
                  f"cảnh báo lần 1: {t['lan1']}, lần 2: {t['lan2']}. Bỏ qua {t['bo_qua_khoa_khac']} dòng của khoa khác. Lý do cảnh báo do admin nhập thủ công.", "ok")
        except Exception as e:
            flash(f"Lỗi nhập file: {e}", "err")
    return redirect(url_for("tinh_trang", hk=hk, nam=nam))


@app.get("/tinh-trang/xuat")
@login_required
def tt_xuat():
    hk, nam = _tt_args()
    kq, kh, khoa, q = (request.args.get(k, "") for k in ("kq", "kh", "khoa", "q"))
    rows = [[i, r["mssv"], r["ho_lot"], r["ten"], r["ma_lop"], r["khoa_hoc"], r["ket_qua"], r["dtb_hk"], r["xl_hk"],
             r["no_hk"], r["dtb_tl"], r["xl_tl"], r["no_tl"], r["ly_do"] or ""]
            for i, r in enumerate(db().ds_tinh_trang(hk, nam, kq, kh, q, khoa), 1)]
    return gui_excel(f"TinhTrang_{hk}_{nam}.xlsx", None,
                     f"TÌNH TRẠNG SINH VIÊN – XÉT TIẾN ĐỘ HỌC TẬP {hk} NĂM HỌC {nam}",
                     ["STT", "MSSV", "Họ lót", "Tên", "Lớp", "Khóa", "Kết quả xét tiến độ học tập", "ĐTB HK",
                      "XL HK", "Nợ HK", "ĐTB TL", "XL TL", "Nợ TL", "Lý do"], rows,
                     [6, 14, 22, 10, 14, 7, 18, 8, 9, 8, 8, 9, 8, 70])


# ------------------------------------------------------------------ nhập lý do bằng Excel / chi tiết sinh viên
def _bo_dau(s):
    s = unicodedata.normalize("NFD", str(s or "")).replace("đ", "d").replace("Đ", "D")
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return " ".join(s.lower().split())


def _o_mssv(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    v = str(v).strip()
    return v[:-2] if v.endswith(".0") and v[:-2].isdigit() else v


def doc_ly_do_excel(fobj):
    """Đọc file Excel có cột MSSV và Lý do (dòng tiêu đề có thể nằm trong 15 dòng đầu).
    Trả về dict {mssv: ly_do}; MSSV lặp thì lấy dòng cuối."""
    from openpyxl import load_workbook
    wb = load_workbook(fobj, read_only=True, data_only=True)
    try:
        ws = wb.active
        c_m = c_l = None
        kq = {}
        for idx, row in enumerate(ws.iter_rows(values_only=True)):
            if c_m is None:
                if idx >= 15:
                    break
                hs = [_bo_dau(x) for x in row]
                m = next((i for i, h in enumerate(hs) if h in ("mssv", "ma sv", "ma sinh vien", "ma so sinh vien")), None)
                l = next((i for i, h in enumerate(hs) if h.startswith("ly do")), None)
                if m is not None and l is not None:
                    c_m, c_l = m, l
                continue
            if len(row) <= max(c_m, c_l):
                continue
            ma = _o_mssv(row[c_m])
            if ma:
                kq[ma] = str(row[c_l] if row[c_l] is not None else "").strip()
        if c_m is None:
            raise ValueError("không tìm thấy dòng tiêu đề có cột MSSV và Lý do")
        return kq
    finally:
        wb.close()


@app.post("/tinh-trang/nhap-ly-do")
@admin_required
def tt_nhap_ly_do():
    hk, nam = _tt_args()
    f = request.files.get("file")
    ghi_de = request.form.get("ghi_de") == "1"
    if not f or not f.filename.lower().endswith(".xlsx"):
        flash("Hãy chọn file .xlsx có cột MSSV và Lý do.", "err")
        return redirect(url_for("tinh_trang", hk=hk, nam=nam))
    try:
        ds = doc_ly_do_excel(io.BytesIO(f.read()))
    except Exception as e:
        flash(f"Lỗi đọc file: {e}", "err")
        return redirect(url_for("tinh_trang", hk=hk, nam=nam))
    con = db().con
    hien_co = {r["mssv"]: (r["ly_do"] or "") for r in
               con.execute("SELECT mssv, ly_do FROM tinh_trang WHERE hoc_ky=? AND nam_hoc=?", (hk, nam))}
    cap_nhat, khong_co, trong, giu = [], [], 0, 0
    for m, ly in ds.items():
        if m not in hien_co:
            khong_co.append(m)
        elif not ly:
            trong += 1
        elif hien_co[m] and not ghi_de:
            giu += 1
        else:
            cap_nhat.append((ly, m, hk, nam))
    con.executemany("UPDATE tinh_trang SET ly_do=? WHERE mssv=? AND hoc_ky=? AND nam_hoc=?", cap_nhat)
    con.commit()
    msg = f"Đã cập nhật lý do cho {len(cap_nhat)} sinh viên ({hk} {nam})."
    if giu:
        msg += f" Giữ nguyên {giu} sinh viên đã có lý do (tích 'Ghi đè' nếu muốn thay)."
    if trong:
        msg += f" Bỏ qua {trong} dòng để trống lý do."
    if khong_co:
        msg += (f" {len(khong_co)} MSSV không có trong danh sách xét tiến độ kỳ này: "
                + ", ".join(khong_co[:10]) + ("…" if len(khong_co) > 10 else "") + ".")
    flash(msg, "ok" if cap_nhat else "err")
    return redirect(url_for("tinh_trang", hk=hk, nam=nam))


@app.get("/tinh-trang/mau-ly-do")
@login_required
def tt_mau_ly_do():
    """File Excel mẫu: các sinh viên đang bị cảnh báo trong kỳ đang xem, kèm lý do hiện có."""
    hk, nam = _tt_args()
    rows = [[r["mssv"], f'{r["ho_lot"]} {r["ten"]}'.strip(), r["ma_lop"], r["ket_qua"], r["ly_do"] or ""]
            for r in db().ds_tinh_trang(hk, nam) if r["ket_qua"] != KET_QUA[0]]
    return gui_excel(f"MauLyDo_{hk}_{nam}.xlsx", None,
                     f"DANH SÁCH SINH VIÊN CẢNH BÁO – NHẬP LÝ DO {hk} NĂM HỌC {nam}",
                     ["MSSV", "Họ tên", "Lớp", "Kết quả xét tiến độ học tập", "Lý do"], rows,
                     [14, 28, 14, 22, 70])


@app.route("/tinh-trang/<mssv>/chi-tiet")
@login_required
def tt_chi_tiet(mssv):
    hk, nam = _tt_args()
    sv = db().con.execute("SELECT * FROM sinh_vien WHERE mssv=?", (mssv,)).fetchone()
    if not sv:
        abort(404)
    ls = [dict(r) for r in db().con.execute(
        "SELECT * FROM tinh_trang WHERE mssv=? ORDER BY nam_hoc DESC, hoc_ky", (mssv,))]
    canh_bao = [r for r in ls if r["ket_qua"] != KET_QUA[0]]
    thong_ke = dict(so_ky=len(ls), canh_bao=len(canh_bao),
                    thieu_ly_do=sum(1 for r in canh_bao if not (r.get("ly_do") or "").strip()))
    drl = [dict(hoc_ky=r["hoc_ky"], nam_hoc=r["nam_hoc"], diem=r["diem"], xep_loai=r["xep_loai"] or "", ly_do=r["ly_do"])
           for r in db().drl_cua_sv(mssv)]
    return render_template("tt_chi_tiet.html", sv=sv, ls=ls, drl=drl, tk=thong_ke, hk=hk, nam=nam)


@app.get("/api/sinh-vien/<mssv>")
@login_required
def api_sinh_vien(mssv):
    """Hồ sơ một sinh viên: thông tin + điểm rèn luyện + lịch sử tình trạng (cho drawer 'Xem')."""
    d = db()
    sv = d.con.execute("SELECT * FROM sinh_vien WHERE mssv=?", (mssv,)).fetchone()
    if not sv:
        return jsonify(error="Không tìm thấy MSSV"), 404
    tt = [dict(r) for r in d.con.execute(
        """SELECT hoc_ky,nam_hoc,ket_qua,dtb_hk,xl_hk,dtb_tl,xl_tl,ly_do FROM tinh_trang WHERE mssv=?
           ORDER BY nam_hoc DESC, CASE hoc_ky WHEN 'HK Hè' THEN 3 WHEN 'HKII' THEN 2 ELSE 1 END DESC""", (mssv,))]
    return jsonify(sinh_vien=dict(sv), tinh_trang=tt, diem_ren_luyen=[dict(r) for r in d.drl_cua_sv(mssv)],
                   so_ky_canh_bao=sum(1 for r in tt if r["ket_qua"] != KET_QUA[0]))


@app.get("/api/kiem-tra-dong-bo")
@login_required
def api_kiem_tra_dong_bo():
    """Báo cáo nhất quán dữ liệu giữa sinh viên / điểm rèn luyện / tình trạng."""
    return jsonify(db().kiem_tra_dong_bo())


routes_chamdiem.dang_ky(app, db, login_required, gui_excel)  # Quản lý chấm điểm / điểm danh


if __name__ == "__main__":
    app.run(host=os.environ.get("HOST", "127.0.0.1"), port=int(os.environ.get("PORT", 5000)), debug=False)