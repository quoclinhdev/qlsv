# -*- coding: utf-8 -*-
"""Lớp CSDL SQLite + nghiệp vụ cho phần mềm quản lý sinh viên."""
import re
import sqlite3
import unicodedata
from datetime import date
from pathlib import Path

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from chamdiem import ChamDiemMixin, SCHEMA as SCHEMA_CHAM_DIEM

import os
DB_PATH = Path(os.environ.get("QLSV_DB", Path(__file__).with_name("quanly_sinhvien.db")))

HOC_KY = ["HKI", "HKII", "HK Hè"]
KET_QUA = ["Bình thường", "Cảnh báo lần 1", "Cảnh báo lần 2"]

# ---------------------------------------------------------------- ĐIỂM RÈN LUYỆN
# Ngưỡng dưới của từng mức (điểm >= ngưỡng thì đạt mức đó), dưới 35 là "Kém"
QUY_TAC_DRL = [(90, "Xuất sắc"), (80, "Tốt"), (65, "Khá"), (50, "TB"), (35, "Yếu")]
# Mã đặc biệt nhập vào cột ĐIỂM
MA_DAC_BIET = {1: "Nghỉ", 2: "Bảo lưu", 3: "Không nộp"}


def xep_loai_drl(diem):
    """Trả về xếp loại theo điểm rèn luyện (hoặc mã 1/2/3)."""
    if diem is None:
        return ""
    if float(diem) in (1.0, 2.0, 3.0):
        return MA_DAC_BIET[int(diem)]
    for nguong, ten in QUY_TAC_DRL:
        if diem >= nguong:
            return ten
    return "Kém"


def parse_diem(text):
    """Chuỗi -> số điểm hợp lệ. '' -> None. Sai -> ValueError."""
    t = str(text).strip().replace(",", ".")
    if t == "":
        return None
    v = float(t)
    if v in (1.0, 2.0, 3.0):
        return v
    if not 0 <= v <= 100:
        raise ValueError("Điểm phải từ 0 đến 100 (hoặc mã 1=Nghỉ, 2=Bảo lưu, 3=Không nộp)")
    return v


def nam_hoc_mac_dinh():
    t = date.today()
    return f"{t.year}-{t.year + 1}" if t.month >= 8 else f"{t.year - 1}-{t.year}"


def nam_hoc_hop_le(s):
    """'2026-2027' hợp lệ khi năm sau = năm trước + 1."""
    m = re.fullmatch(r"(\d{4})-(\d{4})", str(s or "").strip())
    return bool(m) and int(m.group(2)) == int(m.group(1)) + 1


def vi_key(s):
    s = (s or "").lower().replace("đ", "d\uffff")
    return unicodedata.normalize("NFD", s)


def bo_dau(s):
    s = unicodedata.normalize("NFD", str(s or "").lower().replace("đ", "d"))
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


SCHEMA = """
PRAGMA foreign_keys = ON;
CREATE TABLE IF NOT EXISTS sinh_vien(
    mssv TEXT PRIMARY KEY, ho_lot TEXT, ten TEXT,
    ma_lop TEXT, lop TEXT, khoa_hoc TEXT, nganh TEXT, khoa TEXT);
CREATE TABLE IF NOT EXISTS diem_ren_luyen(
    mssv TEXT NOT NULL, hoc_ky TEXT NOT NULL, nam_hoc TEXT NOT NULL,
    diem REAL, xep_loai TEXT, ly_do TEXT DEFAULT '',
    PRIMARY KEY(mssv, hoc_ky, nam_hoc),
    FOREIGN KEY(mssv) REFERENCES sinh_vien(mssv) ON DELETE CASCADE);
CREATE TABLE IF NOT EXISTS tinh_trang(
    mssv TEXT NOT NULL, hoc_ky TEXT NOT NULL, nam_hoc TEXT NOT NULL,
    ket_qua TEXT DEFAULT 'Bình thường',
    dtb_hk REAL, stc_hk INTEGER, xl_hk TEXT, no_hk INTEGER,
    dtb_tl REAL, stc_tl INTEGER, xl_tl TEXT, no_tl INTEGER,
    ly_do TEXT DEFAULT '',
    PRIMARY KEY(mssv, hoc_ky, nam_hoc),
    FOREIGN KEY(mssv) REFERENCES sinh_vien(mssv) ON DELETE CASCADE);
CREATE TABLE IF NOT EXISTS users(
    username TEXT PRIMARY KEY, pw_hash TEXT NOT NULL, role TEXT NOT NULL DEFAULT 'user');
CREATE INDEX IF NOT EXISTS ix_sv_lop ON sinh_vien(ma_lop);
CREATE INDEX IF NOT EXISTS ix_drl_ky ON diem_ren_luyen(nam_hoc, hoc_ky);
CREATE INDEX IF NOT EXISTS ix_tt_ky ON tinh_trang(nam_hoc, hoc_ky);
-- Nguồn dùng chung: mọi module/ứng dụng khác của dự án chỉ cần đọc view này
DROP VIEW IF EXISTS v_sinh_vien_tong_hop;
CREATE VIEW v_sinh_vien_tong_hop AS
SELECT s.mssv, s.ho_lot, s.ten, s.lop, s.ma_lop, s.khoa_hoc, s.nganh, s.khoa,
       k.hoc_ky, k.nam_hoc,
       d.diem AS drl_diem, d.xep_loai AS drl_xep_loai, d.ly_do AS drl_ly_do,
       t.ket_qua, t.dtb_hk, t.xl_hk, t.dtb_tl, t.xl_tl, t.ly_do
FROM sinh_vien s
JOIN (SELECT mssv, hoc_ky, nam_hoc FROM diem_ren_luyen
      UNION SELECT mssv, hoc_ky, nam_hoc FROM tinh_trang) k ON k.mssv = s.mssv
LEFT JOIN diem_ren_luyen d ON d.mssv=k.mssv AND d.hoc_ky=k.hoc_ky AND d.nam_hoc=k.nam_hoc
LEFT JOIN tinh_trang t ON t.mssv=k.mssv AND t.hoc_ky=k.hoc_ky AND t.nam_hoc=k.nam_hoc;
"""


def _s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _f(v):
    try:
        return float(v) if not isinstance(v, str) or not v.startswith("#") else None
    except (TypeError, ValueError):
        return None


def _i(v):
    f = _f(v)
    return None if f is None else int(f)


def _chunks(seq, n=500):
    seq = [str(x) for x in seq]
    for i in range(0, len(seq), n):
        yield seq[i:i + n]


class DB(ChamDiemMixin):
    def __init__(self, path=DB_PATH):
        self.con = sqlite3.connect(str(path))
        self.con.row_factory = sqlite3.Row
        self._migrate()
        self.con.executescript(SCHEMA + SCHEMA_CHAM_DIEM)

    def _migrate(self):
        """Nâng cấp DB cũ: thêm cột lý do cho điểm rèn luyện (giữ nguyên dữ liệu đang có)."""
        self._migrate_cham_diem()
        cols = [r[1] for r in self.con.execute("PRAGMA table_info(diem_ren_luyen)")]
        if cols and "ly_do" not in cols:
            self.con.execute("ALTER TABLE diem_ren_luyen ADD COLUMN ly_do TEXT DEFAULT ''")
            self.con.commit()

    # ---------- sinh viên
    def luu_sinh_vien(self, d):
        self.con.execute(
            """INSERT INTO sinh_vien(mssv,ho_lot,ten,ma_lop,lop,khoa_hoc,nganh,khoa)
               VALUES(:mssv,:ho_lot,:ten,:ma_lop,:lop,:khoa_hoc,:nganh,:khoa)
               ON CONFLICT(mssv) DO UPDATE SET
                 ho_lot=COALESCE(NULLIF(excluded.ho_lot,''),ho_lot),
                 ten=COALESCE(NULLIF(excluded.ten,''),ten),
                 ma_lop=COALESCE(NULLIF(excluded.ma_lop,''),ma_lop),
                 lop=COALESCE(NULLIF(excluded.lop,''),lop),
                 khoa_hoc=COALESCE(NULLIF(excluded.khoa_hoc,''),khoa_hoc),
                 nganh=COALESCE(NULLIF(excluded.nganh,''),nganh),
                 khoa=COALESCE(NULLIF(excluded.khoa,''),khoa)""", d)
        self.con.commit()

    def sua_sinh_vien(self, d):
        """Ghi đè đầy đủ (dùng cho form sửa, cho phép xóa trống ô)."""
        self.con.execute(
            """UPDATE sinh_vien SET ho_lot=:ho_lot,ten=:ten,ma_lop=:ma_lop,lop=:lop,
               khoa_hoc=:khoa_hoc,nganh=:nganh,khoa=:khoa WHERE mssv=:mssv""", d)
        self.con.commit()

    def xoa_sinh_vien(self, mssv):
        self.con.execute("PRAGMA foreign_keys=ON")
        self.con.execute("DELETE FROM sinh_vien WHERE mssv=?", (mssv,))
        self.con.commit()

    def xoa_sinh_vien_nhieu(self, mssvs):
        """Xóa nhiều sinh viên (kèm điểm rèn luyện, tình trạng nhờ ON DELETE CASCADE)."""
        self.con.execute("PRAGMA foreign_keys=ON")
        n = 0
        for ds in _chunks(mssvs):
            n += self.con.execute(
                f"DELETE FROM sinh_vien WHERE mssv IN ({','.join('?' * len(ds))})", ds).rowcount
        self.con.commit()
        return n

    def close(self):
        self.con.close()

    def ds_nam_hoc(self, them=None):
        """Danh sách năm học để chọn: năm đã có dữ liệu + dải năm quanh hiện tại (mới nhất trước)."""
        ys = set()
        for t in ("diem_ren_luyen", "tinh_trang"):
            ys.update(r[0] for r in self.con.execute(f"SELECT DISTINCT nam_hoc FROM {t}"))
        n = int(nam_hoc_mac_dinh()[:4])
        ys.update(f"{y}-{y + 1}" for y in range(n - 5, n + 2))
        if them:
            ys.add(them)
        return sorted((y for y in ys if nam_hoc_hop_le(y)), reverse=True)

    def ds_sinh_vien(self, tu_khoa="", ma_lop=""):
        rows = self.con.execute("SELECT * FROM sinh_vien").fetchall()
        k = bo_dau(tu_khoa)
        out = []
        for r in rows:
            if ma_lop and r["ma_lop"] != ma_lop:
                continue
            if k and k not in bo_dau(f"{r['mssv']} {r['ho_lot']} {r['ten']}"):
                continue
            out.append(r)
        out.sort(key=lambda r: (r["ma_lop"] or "", vi_key(r["ten"]), vi_key(r["ho_lot"])))
        return out

    def ds_lop(self):
        return [r[0] for r in self.con.execute(
            "SELECT DISTINCT ma_lop FROM sinh_vien WHERE ma_lop<>'' ORDER BY ma_lop")]

    def ds_khoa(self):
        return [r[0] for r in self.con.execute(
            "SELECT DISTINCT khoa FROM sinh_vien WHERE khoa<>'' ORDER BY khoa")]

    def ds_khoa_hoc(self):
        return [r[0] for r in self.con.execute(
            "SELECT DISTINCT khoa_hoc FROM sinh_vien WHERE khoa_hoc<>'' ORDER BY khoa_hoc")]

    # ---------- điểm rèn luyện
    def ds_drl(self, hoc_ky, nam_hoc, ma_lop=""):
        rows = self.con.execute(
            """SELECT s.mssv,s.ho_lot,s.ten,s.lop,s.ma_lop,d.diem,d.xep_loai,COALESCE(d.ly_do,'') AS ly_do
               FROM sinh_vien s LEFT JOIN diem_ren_luyen d
                 ON d.mssv=s.mssv AND d.hoc_ky=? AND d.nam_hoc=?
               WHERE (?='' OR s.ma_lop=?)""", (hoc_ky, nam_hoc, ma_lop, ma_lop)).fetchall()
        return sorted(rows, key=lambda r: (r["ma_lop"] or "", vi_key(r["ten"]), vi_key(r["ho_lot"])))

    def luu_drl(self, mssv, hoc_ky, nam_hoc, diem, ly_do=None):
        """Lưu điểm. Lý do chỉ có nghĩa với mã đặc biệt 1/2/3 (Nghỉ/Bảo lưu/Không nộp):
        ly_do=None -> giữ lý do cũ; điểm bình thường -> lý do tự xóa."""
        if diem is None:
            self.con.execute("DELETE FROM diem_ren_luyen WHERE mssv=? AND hoc_ky=? AND nam_hoc=?",
                             (mssv, hoc_ky, nam_hoc))
        else:
            ly = (ly_do or "").strip() if (float(diem) in (1.0, 2.0, 3.0) and ly_do is not None) else None
            self.con.execute(
                """INSERT INTO diem_ren_luyen(mssv,hoc_ky,nam_hoc,diem,xep_loai,ly_do) VALUES(?,?,?,?,?,COALESCE(?,''))
                   ON CONFLICT(mssv,hoc_ky,nam_hoc) DO UPDATE SET diem=excluded.diem,xep_loai=excluded.xep_loai,
                     ly_do=CASE WHEN excluded.diem IN (1,2,3) THEN COALESCE(?,ly_do) ELSE '' END""",
                (mssv, hoc_ky, nam_hoc, diem, xep_loai_drl(diem), ly, ly))
        self.con.commit()

    def ly_do_drl(self, mssv, hoc_ky, nam_hoc):
        r = self.con.execute("SELECT ly_do FROM diem_ren_luyen WHERE mssv=? AND hoc_ky=? AND nam_hoc=?",
                             (mssv, hoc_ky, nam_hoc)).fetchone()
        return (r[0] or "") if r else ""

    def xoa_drl_nhieu(self, mssvs, hoc_ky, nam_hoc):
        """Xóa điểm rèn luyện của các SV trong một học kỳ/năm học (SV vẫn được giữ)."""
        n = 0
        for ds in _chunks(mssvs):
            n += self.con.execute(
                f"DELETE FROM diem_ren_luyen WHERE hoc_ky=? AND nam_hoc=? "
                f"AND mssv IN ({','.join('?' * len(ds))})", [hoc_ky, nam_hoc, *ds]).rowcount
        self.con.commit()
        return n

    def cap_nhat_drl_nhieu(self, mssvs, hoc_ky, nam_hoc, che_do, gia_tri, ly_do=None):
        """Cập nhật điểm nhiều SV trong một giao dịch.
        che_do='cong': cộng/trừ gia_tri vào điểm hiện có (kẹp trong 0–100; bỏ qua SV chưa có điểm,
                       mã đặc biệt 1/2/3, hoặc kết quả trùng mã đặc biệt).
        che_do='dat' : đặt điểm bằng gia_tri cho mọi SV được chọn (kể cả chưa có điểm)."""
        ds = list(dict.fromkeys(str(m).strip() for m in mssvs if str(m).strip()))
        ton_tai, cur = set(), {}
        for ck in _chunks(ds):
            ph = ",".join("?" * len(ck))
            ton_tai.update(r[0] for r in self.con.execute(f"SELECT mssv FROM sinh_vien WHERE mssv IN ({ph})", ck))
            for r in self.con.execute(
                    f"SELECT mssv,diem FROM diem_ren_luyen WHERE hoc_ky=? AND nam_hoc=? AND mssv IN ({ph})",
                    [hoc_ky, nam_hoc, *ck]):
                cur[r[0]] = r[1]
        bo_qua = {"khong_co_sinh_vien": 0, "chua_co_diem": 0, "ma_dac_biet": 0, "trung_ma_dac_biet": 0}
        gioi_han, ghi = 0, []
        for m in ds:
            if m not in ton_tai:
                bo_qua["khong_co_sinh_vien"] += 1
                continue
            if che_do == "dat":
                moi = gia_tri
            else:
                cu = cur.get(m)
                if cu is None:
                    bo_qua["chua_co_diem"] += 1
                    continue
                if float(cu) in (1.0, 2.0, 3.0):
                    bo_qua["ma_dac_biet"] += 1
                    continue
                moi = round(cu + gia_tri, 2)
                if moi > 100 or moi < 0:
                    moi, gioi_han = min(100.0, max(0.0, moi)), gioi_han + 1
                if moi in (1.0, 2.0, 3.0):
                    bo_qua["trung_ma_dac_biet"] += 1
                    continue
            ly = (ly_do or "").strip() if (che_do == "dat" and float(moi) in (1.0, 2.0, 3.0) and ly_do) else None
            ghi.append((m, hoc_ky, nam_hoc, moi, xep_loai_drl(moi), ly, ly))
        self.con.executemany(
            """INSERT INTO diem_ren_luyen(mssv,hoc_ky,nam_hoc,diem,xep_loai,ly_do) VALUES(?,?,?,?,?,COALESCE(?,''))
               ON CONFLICT(mssv,hoc_ky,nam_hoc) DO UPDATE SET diem=excluded.diem,xep_loai=excluded.xep_loai,
                 ly_do=CASE WHEN excluded.diem IN (1,2,3) THEN COALESCE(?,ly_do) ELSE '' END""", ghi)
        self.con.commit()
        return {"cap_nhat": [dict(mssv=g[0], diem=g[3], xep_loai=g[4], ly_do=self.ly_do_drl(g[0], hoc_ky, nam_hoc)) for g in ghi],
                "bo_qua": bo_qua, "gioi_han": gioi_han}

    def nhap_drl_excel(self, path, hoc_ky, nam_hoc):
        """Đọc file Excel có cột MSSV + ĐIỂM, hoặc MSSV + Ghi chú (Nghỉ/Bảo lưu/Không nộp) như file
        'Danh sách sinh viên không có điểm rèn luyện'. Cột 'Lý do' (hoặc 'Trạng thái') được lưu làm lý do.
        Trả về (đã_ghi, không_hợp_lệ, mssv_mã_đặc_biệt_còn_thiếu_lý_do)."""
        ws = load_workbook(path, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        hdr = next((i for i, r in enumerate(rows)
                    if any(bo_dau(c) == "mssv" for c in r if c)), None)
        if hdr is None:
            raise ValueError("Không tìm thấy dòng tiêu đề có cột MSSV")
        h = [bo_dau(c) for c in rows[hdr]]
        c_ms = h.index("mssv")
        c_d = next((i for i, x in enumerate(h) if x.startswith("diem")), None)
        c_gc = next((i for i, x in enumerate(h) if x.startswith("ghi chu")), None)
        c_ld = next((i for i, x in enumerate(h) if x.startswith("ly do") or x.startswith("trang thai")), None)
        if c_d is None and c_gc is None:
            raise ValueError("Không tìm thấy cột ĐIỂM (hoặc cột Ghi chú: Nghỉ / Bảo lưu / Không nộp)")
        ma_chu = {"nghi": "1", "bao luu": "2", "khong nop": "3"}
        cell = lambda r, i: _s(r[i]) if i is not None and i < len(r) else ""
        ok, miss, da_ghi = 0, [], set()
        for r in rows[hdr + 1:]:
            ms = cell(r, c_ms)
            if not ms or bo_dau(ms) == "mssv":      # dòng trống / dòng tiêu đề lặp của mục kế tiếp
                continue
            if not self.con.execute("SELECT 1 FROM sinh_vien WHERE mssv=?", (ms,)).fetchone():
                miss.append(ms)
                continue
            txt = cell(r, c_d)
            if txt == "":
                txt = ma_chu.get(bo_dau(cell(r, c_gc)), "")
            try:
                d = parse_diem(txt)
            except ValueError:
                miss.append(ms)
                continue
            if d is not None:
                ly = cell(r, c_ld)
                self.luu_drl(ms, hoc_ky, nam_hoc, d, ly if ly else None)
                da_ghi.add(ms)
                ok += 1
        thieu = [m for (m,) in self.con.execute(
            "SELECT mssv FROM diem_ren_luyen WHERE hoc_ky=? AND nam_hoc=? AND diem IN (1,2,3) AND COALESCE(ly_do,'')=''",
            (hoc_ky, nam_hoc)) if m in da_ghi]
        return ok, miss, thieu

    # ---------- tình trạng / tiến độ
    def ds_tinh_trang(self, hoc_ky, nam_hoc, ket_qua="", khoa_hoc="", tu_khoa="", khoa=""):
        rows = self.con.execute(
            """SELECT s.mssv,s.ho_lot,s.ten,s.lop,s.ma_lop,s.khoa_hoc,s.khoa,t.*
               FROM tinh_trang t JOIN sinh_vien s ON s.mssv=t.mssv
               WHERE t.hoc_ky=? AND t.nam_hoc=?
                 AND (?='' OR t.ket_qua=?) AND (?='' OR s.khoa_hoc=?) AND (?='' OR s.khoa=?)""",
            (hoc_ky, nam_hoc, ket_qua, ket_qua, khoa_hoc, khoa_hoc, khoa, khoa)).fetchall()
        k = bo_dau(tu_khoa)
        rows = [r for r in rows if not k or k in bo_dau(f"{r['mssv']} {r['ho_lot']} {r['ten']}")]
        return sorted(rows, key=lambda r: (r["ma_lop"] or "", vi_key(r["ten"]), vi_key(r["ho_lot"])))

    def thong_ke_tinh_trang(self, hoc_ky, nam_hoc):
        return self.con.execute(
            """SELECT s.khoa_hoc, t.ket_qua, COUNT(*) n FROM tinh_trang t
               JOIN sinh_vien s ON s.mssv=t.mssv WHERE t.hoc_ky=? AND t.nam_hoc=?
               GROUP BY s.khoa_hoc, t.ket_qua""", (hoc_ky, nam_hoc)).fetchall()

    def luu_tinh_trang(self, d):
        cols = ["mssv", "hoc_ky", "nam_hoc", "ket_qua", "dtb_hk", "stc_hk", "xl_hk", "no_hk",
                "dtb_tl", "stc_tl", "xl_tl", "no_tl", "ly_do"]
        d = {c: d.get(c) for c in cols}
        upd = ",".join(f"{c}=excluded.{c}" for c in cols[3:])
        self.con.execute(
            f"INSERT INTO tinh_trang({','.join(cols)}) VALUES({','.join(':'+c for c in cols)}) "
            f"ON CONFLICT(mssv,hoc_ky,nam_hoc) DO UPDATE SET {upd}", d)
        self.con.commit()

    def xoa_tinh_trang_nhieu(self, mssvs, hoc_ky, nam_hoc):
        n = 0
        for ds in _chunks(mssvs):
            n += self.con.execute(
                f"DELETE FROM tinh_trang WHERE hoc_ky=? AND nam_hoc=? "
                f"AND mssv IN ({','.join('?' * len(ds))})", [hoc_ky, nam_hoc, *ds]).rowcount
        self.con.commit()
        return n

    def cap_nhat_ly_do(self, mssv, hoc_ky, nam_hoc, ly_do):
        self.con.execute("UPDATE tinh_trang SET ly_do=? WHERE mssv=? AND hoc_ky=? AND nam_hoc=?",
                         (ly_do, mssv, hoc_ky, nam_hoc))
        self.con.commit()

    def hoc_ky_gan_nhat(self):
        """Kỳ mới nhất có dữ liệu ở BẤT KỲ module nào (DRL hoặc tình trạng) -> mọi màn hình mở cùng một kỳ."""
        r = self.con.execute(
            """SELECT hoc_ky,nam_hoc FROM (
                   SELECT hoc_ky,nam_hoc FROM tinh_trang UNION SELECT hoc_ky,nam_hoc FROM diem_ren_luyen)
               ORDER BY nam_hoc DESC,
               CASE hoc_ky WHEN 'HK Hè' THEN 3 WHEN 'HKII' THEN 2 ELSE 1 END DESC LIMIT 1""").fetchone()
        return (r[0], r[1]) if r else ("HKI", nam_hoc_mac_dinh())

    def drl_cua_sv(self, mssv):
        """Điểm rèn luyện qua các kỳ của một sinh viên (mới nhất trước)."""
        return self.con.execute(
            """SELECT hoc_ky,nam_hoc,diem,xep_loai,COALESCE(ly_do,'') AS ly_do FROM diem_ren_luyen
               WHERE mssv=? AND diem IS NOT NULL
               ORDER BY nam_hoc DESC, CASE hoc_ky WHEN 'HK Hè' THEN 3 WHEN 'HKII' THEN 2 ELSE 1 END DESC""",
            (mssv,)).fetchall()

    def kiem_tra_dong_bo(self):
        """Báo cáo tính nhất quán giữa 3 module quanh bảng sinh_vien (dùng khi cần rà soát)."""
        q = lambda sql: self.con.execute(sql).fetchone()[0]
        return {
            "foreign_keys_bat": q("PRAGMA foreign_keys") == 1,
            "so_sinh_vien": q("SELECT COUNT(*) FROM sinh_vien"),
            "drl_mo_coi": q("SELECT COUNT(*) FROM diem_ren_luyen d LEFT JOIN sinh_vien s USING(mssv) WHERE s.mssv IS NULL"),
            "tinh_trang_mo_coi": q("SELECT COUNT(*) FROM tinh_trang d LEFT JOIN sinh_vien s USING(mssv) WHERE s.mssv IS NULL"),
            "sv_thieu_ho_ten": q("SELECT COUNT(*) FROM sinh_vien WHERE COALESCE(ten,'')=''"),
            "sv_thieu_ma_lop": q("SELECT COUNT(*) FROM sinh_vien WHERE COALESCE(ma_lop,'')=''"),
            "mssv_sai_dinh_dang": q("SELECT COUNT(*) FROM sinh_vien WHERE mssv<>trim(mssv) OR mssv GLOB '*[^0-9]*'"),
            "xep_loai_drl_lech": q("SELECT COUNT(*) FROM diem_ren_luyen WHERE diem IS NOT NULL AND diem NOT IN (1,2,3) AND ("
                                   "xep_loai IS NOT CASE WHEN diem>=90 THEN 'Xuất sắc' WHEN diem>=80 THEN 'Tốt' "
                                   "WHEN diem>=65 THEN 'Khá' WHEN diem>=50 THEN 'TB' WHEN diem>=35 THEN 'Yếu' ELSE 'Kém' END)"),
            "drl_ma_dac_biet_thieu_ly_do": q("SELECT COUNT(*) FROM diem_ren_luyen WHERE diem IN (1,2,3) AND COALESCE(ly_do,'')=''"),
            "drl_ly_do_thua": q("SELECT COUNT(*) FROM diem_ren_luyen WHERE diem NOT IN (1,2,3) AND COALESCE(ly_do,'')<>''"),
            "ky_co_drl_khong_co_tinh_trang": [tuple(r) for r in self.con.execute(
                """SELECT DISTINCT hoc_ky,nam_hoc FROM diem_ren_luyen EXCEPT
                   SELECT DISTINCT hoc_ky,nam_hoc FROM tinh_trang""")],
            "ky_co_tinh_trang_khong_co_drl": [tuple(r) for r in self.con.execute(
                """SELECT DISTINCT hoc_ky,nam_hoc FROM tinh_trang EXCEPT
                   SELECT DISTINCT hoc_ky,nam_hoc FROM diem_ren_luyen""")],
        }

    # ---------- tài khoản
    def dem_user(self):
        return self.con.execute("SELECT COUNT(*) FROM users").fetchone()[0]

    def them_user(self, username, pw_hash, role):
        self.con.execute("INSERT INTO users(username,pw_hash,role) VALUES(?,?,?)", (username, pw_hash, role))
        self.con.commit()

    def lay_user(self, username):
        return self.con.execute("SELECT * FROM users WHERE username=?", (username,)).fetchone()

    def ds_user(self):
        return self.con.execute("SELECT username,role FROM users ORDER BY username").fetchall()

    def xoa_user(self, username):
        self.con.execute("DELETE FROM users WHERE username=?", (username,))
        self.con.commit()

    def doi_mat_khau(self, username, pw_hash):
        self.con.execute("UPDATE users SET pw_hash=? WHERE username=?", (pw_hash, username))
        self.con.commit()

    # ---------- nhập Excel tiến độ học tập
    def nhap_tien_do_excel(self, path, hoc_ky, nam_hoc):
        """Nhập file Excel có sheet 'XET TIEN DO' (kết quả xét) và các sheet 'DSSV CBHT LAN 1/2'
        (chi tiết KQHT). Trả về dict thống kê."""
        wb = load_workbook(path, data_only=True)
        tk = {"sinh_vien": 0, "binh_thuong": 0, "lan1": 0, "lan2": 0, "bo_qua_khoa_khac": 0}
        khoa_hop_le = set()  # các khoa có trong sheet xét tiến độ; sheet CBHT của cả trường chỉ lấy các khoa này
        tien_do, cbht = [], []
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            ten = bo_dau(ws.title)
            m = re.search(r"lan\s*([12])", ten)
            if m and "cbht" in ten:
                cbht.append((int(m.group(1)), rows))
            elif any(c and "ket qua" in bo_dau(c) for r in rows[:10] for c in r) \
                    and any(c and bo_dau(c) == "mssv" for r in rows[:10] for c in r):
                tien_do.append(rows)

        def lop_ten(nganh, kh):
            return f"{nganh} K{kh}".strip() if kh else nganh

        # 1) Kết quả xét tiến độ cho toàn bộ SV
        for rows in tien_do:
            hdr = next(i for i, r in enumerate(rows) if any(c and bo_dau(c) == "mssv" for c in r))
            h = [bo_dau(c) for c in rows[hdr]]
            col = lambda key: next((i for i, x in enumerate(h) if x.startswith(key)), None)
            c = {k: col(k) for k in ("mssv", "ho lot", "ten", "lop", "khoa", "ket qua", "nganh")}
            c_khoa_don_vi = max((i for i, x in enumerate(h) if x == "khoa"), default=None)
            c_khoa_hoc = next((i for i, x in enumerate(h) if x == "khoa"), None)
            for r in rows[hdr + 1:]:
                ms = _s(r[c["mssv"]])
                if not re.fullmatch(r"\d{6,}", ms):
                    continue
                kq_raw = bo_dau(_s(r[c["ket qua"]]))
                kq = KET_QUA[2] if "2" in kq_raw else KET_QUA[1] if "1" in kq_raw else KET_QUA[0]
                kh = _s(r[c_khoa_hoc])
                nganh = _s(r[c["nganh"]]) if c["nganh"] is not None else ""
                donvi = _s(r[c_khoa_don_vi]) if c_khoa_don_vi is not None else ""
                if donvi:
                    khoa_hop_le.add(bo_dau(donvi))
                self.luu_sinh_vien(dict(mssv=ms, ho_lot=_s(r[c["ho lot"]]), ten=_s(r[c["ten"]]),
                                        ma_lop=_s(r[c["lop"]]), lop=lop_ten(nganh, kh),
                                        khoa_hoc=kh, nganh=nganh, khoa=donvi))
                self.con.execute(
                    """INSERT INTO tinh_trang(mssv,hoc_ky,nam_hoc,ket_qua) VALUES(?,?,?,?)
                       ON CONFLICT(mssv,hoc_ky,nam_hoc) DO UPDATE SET ket_qua=excluded.ket_qua""",
                    (ms, hoc_ky, nam_hoc, kq))
                tk["sinh_vien"] += 1
                tk["binh_thuong" if kq == KET_QUA[0] else "lan1" if kq == KET_QUA[1] else "lan2"] += 1

        # 2) Chi tiết KQHT của SV bị cảnh báo (kèm sinh lý do)
        for lan, rows in sorted(cbht, key=lambda x: x[0]):
            kq = KET_QUA[lan]
            for r in rows:
                if len(r) < 14 or not re.fullmatch(r"\d{6,}", _s(r[1])):
                    continue
                ms, kh = _s(r[1]), _s(r[5])
                donvi = _s(r[16]) if len(r) > 16 else ""
                if khoa_hop_le and bo_dau(donvi) not in khoa_hop_le:
                    tk["bo_qua_khoa_khac"] += 1
                    continue
                self.con.execute(  # SV chưa có (file chỉ có sheet CBHT) -> thêm mới, không ghi đè
                    "INSERT OR IGNORE INTO sinh_vien(mssv,ho_lot,ten,ma_lop,lop,khoa_hoc,nganh,khoa) "
                    "VALUES(?,?,?,?,?,?,?,?)",
                    (ms, _s(r[2]), _s(r[3]), _s(r[4]), f"K{kh}" if kh else "", kh, "", donvi))
                d = dict(mssv=ms, hoc_ky=hoc_ky, nam_hoc=nam_hoc, ket_qua=kq,
                         dtb_hk=_f(r[6]), stc_hk=_i(r[7]), xl_hk=_s(r[8]), no_hk=_i(r[9]),
                         dtb_tl=_f(r[10]), stc_tl=_i(r[11]), xl_tl=_s(r[12]), no_tl=_i(r[13]))
                ghi_chu = _s(r[15]) if len(r) > 15 else ""
                # Lý do do admin nhập; chỉ lấy Ghi chú của file nếu là chữ thật (không phải lỗi #REF!)
                d["ly_do"] = ghi_chu if ghi_chu and not ghi_chu.startswith("#") else ""
                self.luu_tinh_trang(d)
        self.con.commit()
        return tk


# ---------------------------------------------------------------- XUẤT EXCEL
def xuat_excel(path, tieu_de, headers, rows, widths=None, phu_luc=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Danh sách"
    ws.cell(1, 1, tieu_de).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(1, 1).alignment = Alignment(horizontal="center")
    thin = Side(style="thin")
    for j, h in enumerate(headers, 1):
        c = ws.cell(3, j, h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for i, r in enumerate(rows, 4):
        for j, v in enumerate(r, 1):
            c = ws.cell(i, j, v)
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for j, w in enumerate(widths or [], 1):
        ws.column_dimensions[ws.cell(3, j).column_letter].width = w
    if phu_luc:
        c0 = len(headers) + 2
        for i, (a, b) in enumerate(phu_luc, 3):
            ws.cell(i, c0, a).font = Font(bold=(i == 3))
            ws.cell(i, c0 + 1, b).font = Font(bold=(i == 3))
    ws.freeze_panes = "A4"
    wb.save(path)